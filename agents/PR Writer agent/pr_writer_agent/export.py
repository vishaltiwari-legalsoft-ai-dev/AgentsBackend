"""Export one version of a PR run as Markdown or .docx (blueprint §10 order).

Input is the wire view the service already builds (``PrRun`` + ``PrVersion``
from ``docs/pr-writer-api.md``), so the file says exactly what the screen
says: the same blocks, the same open markers, the same status. A version that
is not approved carries "DRAFT — NOT APPROVED FOR RELEASE" on its cover.
Nothing here sends anything anywhere; the bytes go back to the caller's
browser.
"""
from __future__ import annotations

import io
import re
from typing import Any

NOT_APPROVED = "DRAFT — NOT APPROVED FOR RELEASE"


def filename(run: dict[str, Any], version: dict[str, Any], ext: str, today: str) -> str:
    def slug(s: str, n: int = 40) -> str:
        return re.sub(r"[^A-Za-z0-9]+", "", str(s or "").title())[:n] or "Draft"

    return (f"{slug(run.get('brand_name'))}_{slug(run.get('content_type_label') or 'PR')}_"
            f"{slug(run.get('news'), 32)}_v{version['version']}_{today}.{ext}")


def _status_line(run: dict[str, Any], version: dict[str, Any]) -> str:
    if run.get("status") == "approved" and run.get("current_version") == version["version"]:
        return f"Approved v{version['version']}"
    if version.get("locked"):
        return f"In review v{version['version']}"
    return f"Draft v{version['version']}"


def _sections(run: dict[str, Any], version: dict[str, Any]) -> list[tuple[str, list[Any]]]:
    """``[(heading, parts)]`` where a part is a paragraph string or a table
    ``("table", header, rows)``."""
    approved = run.get("status") == "approved" and run.get("current_version") == version["version"]
    open_verify = sum(1 for m in version.get("markers", []) if m["kind"] == "verify" and not m["resolved"])
    pending = sum(1 for q in version.get("quotes", []) if q["approval"]["state"] != "approved")
    cover = [
        f"Brand: {run.get('brand_name')}",
        f"Content type: {run.get('content_type_label') or '—'}",
        f"Version: v{version['version']} · {version.get('created_at', '')[:10]}",
        f"Status: {_status_line(run, version)}",
        f"Open [VERIFY] items: {open_verify} · Quotes awaiting approval: {pending}",
        "PR Writer drafts only. It never sends, posts or pitches anything.",
    ]
    if not approved:
        cover.insert(0, NOT_APPROVED)
    heads = [f"{h['text']} ({h['chars']} characters){' — recommended' if h.get('recommended') else ''}"
             for h in version.get("headlines", [])]
    sub = version.get("subheadline") or {}
    out: list[tuple[str, list[Any]]] = [
        ("Cover sheet", cover),
        ("News angle", [version.get("news_angle") or "—"]),
        ("Headline options", heads or ["—"]),
        ("Subheadline", [f"Recommended: {sub.get('recommended') or '—'}", f"Alternate: {sub.get('alternate') or '—'}"]),
        ("Full draft", [b["text"] for b in version.get("blocks", [])]),
        ("Quote approval log", [("table", ["Quote", "Speaker", "Title", "Status", "Evidence"], [
            [q["text"], q["speaker_name"], q["speaker_title"], q["approval"]["state"],
             q["approval"].get("evidence") or ""] for q in version.get("quotes", [])])]),
        ("Fact-check table", [("table", ["#", "Claim as written", "Source", "Date checked", "Status", "Owner to confirm"], [
            [str(c["n"]), c["text"], c["source"]["label"], c.get("date_checked") or "", c["status"],
             c.get("owner_to_confirm") or ""] for c in version.get("claims", [])])]),
    ]
    pkg = run.get("package") or {}
    seo = pkg.get("seo")
    out.append(("SEO", [
        f"Primary keyword: {seo.get('primary_keyword')}",
        f"Secondary keywords: {', '.join(seo.get('secondary_keywords') or [])}",
        f"SEO title: {seo.get('seo_title')}", f"Meta description: {seo.get('meta_description')}",
        f"Slug: {seo.get('slug')}",
    ] if seo else ["SEO wasn't generated."]))
    dist = pkg.get("distribution")
    out.append(("Distribution plan", [f"{c.get('channel')}: {c.get('why')} ({c.get('timing')})"
                                      for c in (dist or {}).get("channels", [])] or ["Distribution plan wasn't generated."]))
    snip = pkg.get("snippets")
    out.append(("Social and pitch snippets", [
        f"LinkedIn: {snip.get('linkedin_post')}", f"Social: {snip.get('social_summary')}",
        f"Reporter pitch: {snip.get('reporter_pitch')}",
    ] if snip else ["Snippets weren't generated."]))
    notes = pkg.get("reviewer_notes")
    out.append(("Reviewer notes", [f"{n['kind'].replace('_', ' ')}: {n['text']}" for n in notes or []]
                or ["Reviewer notes weren't generated."]))
    return out


def to_markdown(run: dict[str, Any], version: dict[str, Any]) -> str:
    lines: list[str] = []
    for heading, parts in _sections(run, version):
        lines.append(f"## {heading}\n")
        for part in parts:
            if isinstance(part, tuple):
                _, header, rows = part
                lines.append("| " + " | ".join(header) + " |")
                lines.append("|" + "---|" * len(header))
                for row in rows:
                    lines.append("| " + " | ".join(str(c).replace("|", "\\|").replace("\n", " ") for c in row) + " |")
                lines.append("")
            else:
                lines.append(str(part) + "\n")
    return "\n".join(lines)


def to_docx(run: dict[str, Any], version: dict[str, Any]) -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading(f"{run.get('brand_name')} — v{version['version']}", level=0)
    for heading, parts in _sections(run, version):
        doc.add_heading(heading, level=1)
        for part in parts:
            if isinstance(part, tuple):
                _, header, rows = part
                table = doc.add_table(rows=1, cols=len(header))
                table.style = "Table Grid"
                for cell, text in zip(table.rows[0].cells, header):
                    cell.text = text
                for row in rows:
                    cells = table.add_row().cells
                    for cell, text in zip(cells, row):
                        cell.text = str(text)
            else:
                doc.add_paragraph(str(part))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
